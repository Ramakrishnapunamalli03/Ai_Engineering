"""
Universal document loader — handles PDF, DOCX, TXT, HTML, and Markdown.
Run: python document_loader.py your_file.pdf
"""
import sys
from pathlib import Path

def load_document(file_path: str) -> str:
    """Load text from any supported document format. Returns a single string."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")

    if suffix == ".pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(file_path)
        chunks = []
        for page in pdf:
            textpage = page.get_textpage()
            chunks.append(textpage.get_text_range() or "")
            textpage.close()
            page.close()
        return "\n".join(chunks)

    if suffix == ".docx":
        from docx import Document
        doc = Document(file_path)
        return "\n".join(para.text for para in doc.paragraphs)

    if suffix == ".html":
        from bs4 import BeautifulSoup
        html = path.read_text(encoding="utf-8")
        soup = BeautifulSoup(html, "html.parser")
        return soup.get_text(separator="\n", strip=True)

    raise ValueError(f"Unsupported format: {suffix}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python document_loader.py <file_path>")
        sys.exit(1)
    text = load_document(sys.argv[1])
    print(f"Loaded {len(text)} characters from {sys.argv[1]}")
    print(f"Preview: {text[:500]}...")